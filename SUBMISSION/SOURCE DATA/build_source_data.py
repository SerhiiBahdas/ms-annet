#!/usr/bin/env python3
"""Build the source-data workbooks (Supplementary Data 1 to 11) of the manuscript figures.

Every figure that shows data has an archive FIGn.zip next to its notebook FIGn.ipynb.
This script copies the tables that a figure draws, and the statistics tables that its
legend quotes, from the archive into one Excel workbook per figure:

    Supplementary Data 1.xlsx ... Supplementary Data 8.xlsx      (Figs 2 to 9)
    Supplementary Data 9.xlsx ... Supplementary Data 11.xlsx     (Supplementary Figs 1 to 3)

Values are copied without rounding, recomputation or reordering. The first sheet of
each workbook (README) lists the sheets with their panels and sample sizes, and every
column with its unit. Each workbook is read back after writing and compared with the
archive, table by table.

Usage:
    python build_source_data.py [OUTPUT_DIR] [--archives ARCHIVE_DIR]

OUTPUT_DIR defaults to the folder of this script and ARCHIVE_DIR to
../FIGURES/DATA AND SCRIPTS relative to it.
"""

import argparse
import hashlib
import io
import json
import math
import sys
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.cell import _writer as cell_writer
from openpyxl.styles import Alignment, Border, Font

NONE = "none"          # unit of labels, counts of categories and dimensionless numbers
COUNT = "count"
EXCEL_MAX_ROWS = 1_048_576


@dataclass
class Sheet:
    name: str        # sheet name in the workbook (at most 31 characters)
    member: str      # table in the archive
    panel: str       # panel or panels that the table belongs to
    content: str     # what the sheet holds and which rows or columns the figure draws
    n: str           # sample size
    shape: tuple     # expected (rows, columns) of the table
    columns: dict    # column -> (description, unit)


@dataclass
class Figure:
    label: str       # "Fig. 2" or "Supplementary Fig. 1"
    archive: str     # archive file name
    title: str       # title of the figure in the manuscript
    sheets: list
    note: str = ""   # optional remark for the README sheet

    @property
    def name(self):
        return f"Supplementary Data {FIGURES.index(self) + 1}"

    @property
    def workbook(self):
        return f"{self.name}.xlsx"


# ----------------------------------------------------------------------------------
# Column descriptions shared by several tables
# ----------------------------------------------------------------------------------
JOINT = {
    "joint": ("Joint, J1 for joint 1 and J2 for joint 2", NONE),
    "joint_label": ("Name of the joint", NONE),
}
TRAJECTORY = {"trajectory": ("Index of the test trajectory, 0 to 49", NONE)}
QUANTITY = {
    "quantity": ("Quantity of the forward simulation that the error refers to", NONE),
    "unit": ("Unit of the error of that quantity", NONE),
}


def primary_test(group_a, group_b, family):
    """Columns of the two-group comparison used throughout the statistics tables."""
    return {
        "both_normal": ("True when both Shapiro-Wilk p values exceed 0.05", NONE),
        "test_name": ("Primary test, one-way ANOVA when both groups are normal and the Kruskal-Wallis "
                      "H-test otherwise", NONE),
        "statistic": (f"Statistic of the primary test (F or H), {group_a} against {group_b}", NONE),
        "p_value": ("p value of the primary test", NONE),
        "fdr_q_value": (f"Benjamini-Hochberg adjusted p value of the primary test, {family}", NONE),
        "significant_fdr_0_05": ("True when the adjusted p value of the primary test is below 0.05", NONE),
    }


def paired_test(group_a, group_b):
    return {
        "wilcoxon_paired_statistic": ("Statistic of the two-sided Wilcoxon signed-rank test paired by trajectory, "
                                      f"{group_a} against {group_b}", NONE),
        "wilcoxon_paired_p_two_sided": ("p value of the paired Wilcoxon signed-rank test", NONE),
    }


def trace_columns(sources, quantities, drawn):
    """Columns <source>_<quantity>_j<joint> of an example trajectory."""
    columns = {"time_s": ("Time (x axis)", "s")}
    for source, source_text in sources.items():
        for quantity, (quantity_text, unit) in quantities.items():
            for joint in (1, 2):
                name = f"{source}_{quantity}_j{joint}"
                mark = ", drawn" if (source, quantity) in drawn else ", not drawn"
                columns[name] = (f"{quantity_text} of joint {joint}, {source_text}{mark}", unit)
    return columns


# ----------------------------------------------------------------------------------
# Fig. 2
# ----------------------------------------------------------------------------------
FIG2 = Figure(
    label="Fig. 2", archive="FIG2.zip",
    title="Systematic hyperparameter optimization",
    sheets=[
        Sheet(
            name="Fig. 2a-g", member="fig2_hyperparameter_trials.csv", panel="a-g",
            content="One row per trial of the hyperparameter search. Each panel plots torque_rmse_nm against one "
                    "hyperparameter (a learning_rate, b network_depth, c nodes_layer1, d activation, e nodes_layer2, "
                    "f nodes_layer3, g nodes_layer4).",
            n="n = 500 trials, one network configuration each. Hidden layers 2, 3 and 4 exist in 481, 404 and "
              "345 trials.",
            shape=(500, 11),
            columns={
                "trial_number": ("Index of the trial in the search", NONE),
                "trial_state": ("Final state of the trial, COMPLETE or PRUNED", NONE),
                "mse_loss": ("Objective value of the trial, the validation mean squared torque error", "Nm^2"),
                "torque_rmse_nm": ("Validation torque RMSE, the square root of mse_loss (y axis)", "Nm"),
                "activation": ("Activation function (symbol shape, x axis of d)", NONE),
                "learning_rate": ("Learning rate (x axis of a)", NONE),
                "network_depth": ("Number of hidden layers (x axis of b)", COUNT),
                "nodes_layer1": ("Nodes in hidden layer 1 (x axis of c)", COUNT),
                "nodes_layer2": ("Nodes in hidden layer 2 (x axis of e), empty when the layer is absent", COUNT),
                "nodes_layer3": ("Nodes in hidden layer 3 (x axis of f), empty when the layer is absent", COUNT),
                "nodes_layer4": ("Nodes in hidden layer 4 (x axis of g), empty when the layer is absent", COUNT),
            },
        ),
        Sheet(
            name="Fig. 2a-g best", member="fig2_best_configuration.json", panel="a-g",
            content="Best configuration of the search, drawn as the red diamond in every panel.",
            n="n = 1 configuration (trial 220)",
            shape=(1, 10),
            columns={
                "learning_rate": ("Learning rate", NONE),
                "network_depth": ("Number of hidden layers", COUNT),
                "activation": ("Activation function", NONE),
                "nodes_layer1": ("Nodes in hidden layer 1", COUNT),
                "nodes_layer2": ("Nodes in hidden layer 2", COUNT),
                "nodes_layer3": ("Nodes in hidden layer 3", COUNT),
                "nodes_layer4": ("Nodes in hidden layer 4", COUNT),
                "best_trial_number": ("Index of the trial in the search", NONE),
                "best_mse_loss": ("Validation mean squared torque error of the trial", "Nm^2"),
                "best_torque_rmse_nm": ("Validation torque RMSE of the trial (y value of the diamond)", "Nm"),
            },
        ),
        Sheet(
            name="Fig. 2h", member="fig2_hyperparameter_sensitivity.csv", panel="h",
            content="Relative sensitivity of the search to each hyperparameter (bar lengths).",
            n="n = 4 hyperparameters",
            shape=(4, 4),
            columns={
                "parameter_key": ("Name of the hyperparameter in the search", NONE),
                "parameter_label": ("Label of the bar", NONE),
                "sensitivity_fraction": ("Importance of the hyperparameter (fANOVA) as a fraction of the total "
                                         "(bar length)", NONE),
                "sensitivity_percent": ("The same importance in percent (number beside the bar)", "%"),
            },
        ),
    ],
)

# ----------------------------------------------------------------------------------
# Fig. 3
# ----------------------------------------------------------------------------------
DATASET_SIZE = {"dataset_size": ("Number of training samples of the network", COUNT)}

FIG3 = Figure(
    label="Fig. 3", archive="FIG3.zip",
    title="Dataset-size effects on inverse-dynamics accuracy and training cost",
    sheets=[
        Sheet(
            name="Fig. 3a", member="fig3_test_rmse_history.csv", panel="a",
            content="Validation RMSE during training (lines), evaluated every 100 epochs.",
            n="n = 9 networks (one per training-set size), 30 evaluations each",
            shape=(270, 4),
            columns={
                **DATASET_SIZE,
                "epoch": ("Training epoch of the evaluation (x axis)", NONE),
                "test_mse": ("Mean squared torque error on the validation split", "Nm^2"),
                "test_rmse_nm": ("Torque RMSE on the validation split (y axis)", "Nm"),
            },
        ),
        Sheet(
            name="Fig. 3a checkpoints", member="fig3_best_checkpoints.csv", panel="a",
            content="Checkpoint with the lowest validation RMSE of each network (markers).",
            n="n = 9 networks",
            shape=(9, 4),
            columns={
                **DATASET_SIZE,
                "best_epoch": ("Epoch of the lowest validation RMSE (x value of the marker)", NONE),
                "best_test_rmse_nm": ("Lowest validation torque RMSE (y value of the marker)", "Nm"),
                "best_test_mse": ("Mean squared torque error at that epoch", "Nm^2"),
            },
        ),
        Sheet(
            name="Fig. 3b-c", member="fig3_pairwise_statistics.csv", panel="b, c",
            content="Pairwise comparisons of the torque RMSE distributions across training-set sizes, joint 1 in b "
                    "and joint 2 in c. A chord is grey when significant_fdr_0_05 is True and coloured when it is "
                    "False. The compared values are on the sheet Fig. 3d upper.",
            n="n = 50 test trajectories per training-set size, 36 comparisons per joint",
            shape=(72, 9),
            columns={
                "joint": JOINT["joint"],
                "dataset_size_a": ("Training-set size of the first network of the pair", COUNT),
                "dataset_size_b": ("Training-set size of the second network of the pair", COUNT),
                "n_a": ("Number of test trajectories of the first network", COUNT),
                "n_b": ("Number of test trajectories of the second network", COUNT),
                "mann_whitney_u": ("U statistic of the two-sided Mann-Whitney U test", NONE),
                "raw_p_value": ("p value of that test", NONE),
                "fdr_adjusted_p": ("Benjamini-Hochberg adjusted p value across the 36 comparisons of the joint", NONE),
                "significant_fdr_0_05": ("True when the adjusted p value is below 0.05", NONE),
            },
        ),
        Sheet(
            name="Fig. 3d upper", member="fig3_rmse_values.csv", panel="d, upper panel",
            content="Torque RMSE of every test trajectory (points). Boxes, centre lines and whiskers are the "
                    "quartiles, medians and 5th and 95th percentiles of these values.",
            n="n = 50 test trajectories per training-set size and joint",
            shape=(900, 5),
            columns={
                **DATASET_SIZE,
                "joint": JOINT["joint"],
                "trajectory_index": ("Index of the test trajectory, 0 to 49", NONE),
                "torque_rmse_nm": ("Torque RMSE of the trajectory (y axis)", "Nm"),
                "source_result_key": ("Sampling step of the test trajectories, the key of the result set in the "
                                      "stored results", "s"),
            },
        ),
        Sheet(
            name="Fig. 3d fits", member="fig3_power_law_fits.csv", panel="d, upper panel",
            content="Power-law fits to the median RMSE of each training-set size (solid lines), "
                    "RMSE(n) = a n^(-b) + c with n the number of training samples.",
            n="n = 9 medians per joint",
            shape=(2, 5),
            columns={
                "joint": JOINT["joint"],
                "a": ("Scale of the fit", "Nm"),
                "b": ("Exponent of the fit", NONE),
                "c": ("Offset of the fit", "Nm"),
                "r_squared": ("Coefficient of determination of the fit", NONE),
            },
        ),
        Sheet(
            name="Fig. 3d lower", member="fig3_training_times.csv", panel="d, lower panel",
            content="Wall-clock training time of each network.",
            n="n = 9 networks, one training run each",
            shape=(9, 4),
            columns={
                **DATASET_SIZE,
                "elapsed_seconds": ("Training time (y axis)", "s"),
                "elapsed_minutes": ("Training time", "min"),
                "elapsed_hours": ("Training time", "h"),
            },
        ),
    ],
)

# ----------------------------------------------------------------------------------
# Fig. 4
# ----------------------------------------------------------------------------------
FIG4 = Figure(
    label="Fig. 4", archive="FIG4.zip",
    title="ANNet achieves a lower median torque error than an otherwise identical direct regressor",
    sheets=[
        Sheet(
            name="Fig. 4a-b", member="fig4_rmse_values.csv", panel="a, b",
            content="Torque RMSE of every test trajectory for ANNet and for the direct regressor (MLP), joint 1 in a "
                    "and joint 2 in b. A line joins the two rows that share joint and trajectory.",
            n="n = 50 test trajectories per network and joint, each evaluated by both networks",
            shape=(200, 4),
            columns={
                "model": ("Network, ANNet or the direct torque regressor (MLP)", NONE),
                "joint": JOINT["joint"],
                **TRAJECTORY,
                "rmse_nm": ("Torque RMSE of the trajectory (y axis)", "Nm"),
            },
        ),
        Sheet(
            name="Fig. 4a-b statistics", member="fig4_statistics.csv", panel="a, b",
            content="Comparison of the two networks for each joint (asterisks and the tests of the legend).",
            n="n = 50 trajectories per group, 2 comparisons",
            shape=(2, 26),
            columns={
                **JOINT,
                "n_annet": ("Number of trajectories of ANNet", COUNT),
                "n_mlp": ("Number of trajectories of the direct regressor", COUNT),
                "shapiro_p_annet": ("Shapiro-Wilk p value of the ANNet errors", NONE),
                "shapiro_p_mlp": ("Shapiro-Wilk p value of the errors of the direct regressor", NONE),
                **primary_test("direct regressor", "ANNet", "across the 2 joints"),
                "mann_whitney_u_mlp_vs_annet": ("U statistic of the two-sided Mann-Whitney U test, direct regressor "
                                                "against ANNet", NONE),
                "mann_whitney_p_two_sided": ("p value of the Mann-Whitney U test", NONE),
                **paired_test("direct regressor", "ANNet"),
                "mean_annet_nm": ("Mean ANNet error", "Nm"),
                "sd_annet_nm": ("Standard deviation of the ANNet errors", "Nm"),
                "median_annet_nm": ("Median ANNet error (horizontal bar)", "Nm"),
                "iqr_annet_nm": ("Interquartile range of the ANNet errors", "Nm"),
                "mean_mlp_nm": ("Mean error of the direct regressor", "Nm"),
                "sd_mlp_nm": ("Standard deviation of the errors of the direct regressor", "Nm"),
                "median_mlp_nm": ("Median error of the direct regressor (horizontal bar)", "Nm"),
                "iqr_mlp_nm": ("Interquartile range of the errors of the direct regressor", "Nm"),
                "median_ratio_mlp_over_annet": ("Median error of the direct regressor divided by that of ANNet", NONE),
                "cliffs_delta_mlp_vs_annet": ("Cliff's delta, direct regressor against ANNet", NONE),
            },
        ),
    ],
)

# ----------------------------------------------------------------------------------
# Fig. 5
# ----------------------------------------------------------------------------------
GRID_COLUMNS = {
    "group_key": ("Key of the set of configurations", NONE),
    "group_label": ("Name of the set of configurations", NONE),
    "label": ("Identifier of the configuration", NONE),
    "point_label": ("Configuration written with N, B and A", NONE),
    "candidate_count": ("Number of candidate acceleration vectors per state update, N", COUNT),
    "bound": ("Bound of the uniform acceleration samples, B (colour)", "rad/s^2"),
    "adam_steps": ("Number of Adam refinement steps, A", COUNT),
    "sample_count": ("Number of state updates evaluated, 50 trajectories of 200 updates", COUNT),
    "random_uniform_vectors": ("Candidates drawn at random, N minus the zero and the previous acceleration", COUNT),
    "mean_abs_angle_error_rad": ("Mean absolute angle error, averaged across the two joints (y axis)", "rad"),
    "mean_abs_angle_error_deg": ("The same error in degrees", "deg"),
    "mean_abs_angle_error_joint1_rad": ("Mean absolute angle error of joint 1", "rad"),
    "mean_abs_angle_error_joint2_rad": ("Mean absolute angle error of joint 2", "rad"),
    "rmse_q": ("RMSE of the simulated angles over both joints", "rad"),
    "rmse_dq": ("RMSE of the simulated velocities over both joints", "rad/s"),
    "rmse_ddq": ("RMSE of the simulated accelerations over both joints", "rad/s^2"),
    "mean_total_ms": ("Mean time of one solver call", "ms"),
    "median_total_ms": ("Median time of one solver call", "ms"),
    "p95_total_ms": ("95th percentile of the time of one solver call", "ms"),
    "mean_forward_ms": ("Mean time of the network evaluations within one solver call", "ms"),
    "median_forward_ms": ("Median time of the network evaluations within one solver call", "ms"),
    "p95_forward_ms": ("95th percentile of the time of the network evaluations within one solver call", "ms"),
    "rollout_wall_s": ("Wall-clock time of the 50 simulated trajectories", "s"),
    "order_idx": ("Position of the configuration in the grid", NONE),
    "sampled_acceleration_vectors": ("Number of sampled acceleration vectors, N (equal to candidate_count)", COUNT),
    "mean_solver_call_ms": ("Mean time of one solver call (equal to mean_total_ms)", "ms"),
    "mean_network_forward_ms": ("Mean time of the network evaluations (equal to mean_forward_ms)", "ms"),
    "mean_update_ms": ("Mean execution time of one state update, rollout_wall_s divided by sample_count "
                       "(x axis)", "ms"),
    "real_time_factor": ("Real-time step of 2 ms divided by mean_update_ms", NONE),
    "runs_faster_than_real_time": ("True when mean_update_ms is below the real-time limit of 2 ms", NONE),
    "time_angle_error_product": ("mean_update_ms multiplied by mean_abs_angle_error_rad", "ms rad"),
}

FIG5 = Figure(
    label="Fig. 5", archive="FIG5.zip",
    title="Accuracy-latency trade-offs of the optimization-based forward dynamics solver",
    sheets=[
        Sheet(
            name="Fig. 5", member="fig5_optimizer_grid_search.csv", panel="single panel",
            content="One row per solver configuration (points). The figure plots mean_abs_angle_error_rad against "
                    "mean_update_ms and colours each point by bound.",
            n="n = 336 solver configurations, each evaluated over 50 trajectories (10,000 state updates)",
            shape=(336, 31),
            columns=GRID_COLUMNS,
        ),
        Sheet(
            name="Fig. 5 Pareto front", member="fig5_pareto_front.csv", panel="single panel",
            content="Configurations on the empirical Pareto frontier (black dashed curve and open circles), a "
                    "subset of the rows of the sheet Fig. 5.",
            n="n = 34 configurations",
            shape=(34, 32),
            columns={"pareto_rank_by_time": ("Position along the frontier, ordered by execution time", NONE),
                     **GRID_COLUMNS},
        ),
        Sheet(
            name="Fig. 5 callouts", member="fig5_highlighted_operating_points.csv", panel="single panel",
            content="The four configurations marked by numbered callouts, a subset of the rows of the sheet Fig. 5.",
            n="n = 4 configurations",
            shape=(4, 33),
            columns={"callout_number": ("Number of the callout in the figure", NONE),
                     "selection": ("Criterion by which the configuration was selected", NONE),
                     **GRID_COLUMNS},
        ),
    ],
)

# ----------------------------------------------------------------------------------
# Fig. 6
# ----------------------------------------------------------------------------------
ROLLOUT_SOURCES = {"numerical": "numerical ODE solution (dashed line)",
                   "neural": "optimization-based neural forward dynamics solver (solid line)"}
ROLLOUT_QUANTITIES = {"position": ("Angle", "rad", "rad"), "velocity": ("Velocity", "rad/s", "rad_s"),
                      "acceleration": ("Acceleration", "rad/s^2", "rad_s2")}
ROLLOUT_COLUMNS = {"time_index": ("Index of the time point", NONE), "time_s": ("Time (x axis)", "s")}
for _quantity, (_text, _unit, _suffix) in ROLLOUT_QUANTITIES.items():
    for _source, _source_text in ROLLOUT_SOURCES.items():
        for _joint in (1, 2):
            ROLLOUT_COLUMNS[f"{_source}_{_quantity}_joint{_joint}_{_suffix}"] = (
                f"{_text} of joint {_joint}, {_source_text}", _unit)

FIG6 = Figure(
    label="Fig. 6", archive="FIG6.zip",
    title="Learned acceleration energy landscape enables accurate and fast forward simulation",
    sheets=[
        Sheet(
            name="Fig. 6a", member="fig6_demo_trajectory.csv", panel="a",
            content="Representative closed-loop rollout at every time point. The figure draws every fourth time "
                    "point of these traces.",
            n="n = 1 trajectory (index 10 of the 50 simulations), 2,500 time points",
            shape=(2500, 14),
            columns=ROLLOUT_COLUMNS,
        ),
        Sheet(
            name="Fig. 6b", member="fig6_rmse_values.csv", panel="b",
            content="RMSE of angle, velocity and acceleration of every simulated trajectory (points). Boxes span "
                    "the interquartile range of these values, centre lines are medians and whiskers extend to "
                    "1.5 times the interquartile range. The open circle marks trial_index 10, the trajectory of a.",
            n="n = 50 independent simulations per joint",
            shape=(100, 6),
            columns={
                "trial_index": ("Index of the simulated trajectory, 0 to 49", NONE),
                "joint_index": ("Joint, 1 or 2", NONE),
                "joint_label": ("Name of the joint", NONE),
                "position_rmse_rad": ("RMSE of the simulated angle against the numerical ODE solution", "rad"),
                "velocity_rmse_rad_s": ("RMSE of the simulated velocity against the numerical ODE solution", "rad/s"),
                "acceleration_rmse_rad_s2": ("RMSE of the simulated acceleration against the numerical ODE solution",
                                             "rad/s^2"),
            },
        ),
        Sheet(
            name="Fig. 6c", member="fig6_trajectory_timing.csv", panel="c",
            content="Computation time of every 5-s trajectory (points). Bars are the medians of these values.",
            n="n = 50 trajectories per solver",
            shape=(50, 3),
            columns={
                "trial_index": ("Index of the simulated trajectory, 0 to 49", NONE),
                "ode_solver_s": ("Computation time of the trajectory with the numerical ODE solver", "s"),
                "neural_solver_s": ("Computation time of the trajectory with the optimization-based neural forward "
                                    "dynamics solver", "s"),
            },
        ),
        Sheet(
            name="Fig. 6d", member="fig6_latency_values.csv", panel="d",
            content="Computation time of every query. The histograms count these values in 40 logarithmic bins per "
                    "decade between 0.02 and 200 ms.",
            n="n = 125,000 forward-solver updates and 10,000 inverse-model torque predictions",
            shape=(135000, 3),
            columns={
                "measurement": ("Type of query, forward_neural_solver for one forward-solver update and "
                                "inverse_model_16k for one torque prediction of the inverse model", NONE),
                "query_index": ("Index of the query within its type", NONE),
                "latency_ms": ("Computation time of the query (x axis)", "ms"),
            },
        ),
        Sheet(
            name="Fig. 6d summary", member="fig6_latency_summary.csv", panel="d",
            content="Summary of the two distributions of the sheet Fig. 6d.",
            n="n = 125,000 forward-solver updates and 10,000 inverse-model torque predictions",
            shape=(2, 8),
            columns={
                "measurement": ("Type of query, as on the sheet Fig. 6d", NONE),
                "n": ("Number of queries", COUNT),
                "min_ms": ("Shortest computation time", "ms"),
                "q1_ms": ("First quartile of the computation time", "ms"),
                "median_ms": ("Median computation time", "ms"),
                "q3_ms": ("Third quartile of the computation time", "ms"),
                "p99_ms": ("99th percentile of the computation time", "ms"),
                "max_ms": ("Longest computation time", "ms"),
            },
        ),
    ],
)

# ----------------------------------------------------------------------------------
# Fig. 7
# ----------------------------------------------------------------------------------
NOISE_PERCENT_INPUT = ("Noise level, the s.d. of the noise in percent of the s.d. of each input over the trajectory",
                       "%")

FIG7 = Figure(
    label="Fig. 7", archive="FIG7.zip",
    title="ANNet reproduces the sensitivity of analytical inverse dynamics to measurement noise",
    sheets=[
        Sheet(
            name="Fig. 7a", member="fig7_example_traces.csv", panel="a",
            content="Torques of the example trajectory with 5% noise on all inputs at every time point. The traces "
                    "show the noise-free and the ANNet torques. The frames show the final 21 samples (2 ms) of all "
                    "three torques.",
            n="n = 1 trajectory (index 11 of the 50 test trajectories), 50,001 time points",
            shape=(50001, 8),
            columns={
                "sample": ("Index of the time point", NONE),
                "time_s": ("Time (x axis)", "s"),
                "noise_free_j1_nm": ("Noise-free torque of joint 1 (black line)", "Nm"),
                "annet_noisy_j1_nm": ("ANNet torque of joint 1 for the noisy inputs (coloured trace)", "Nm"),
                "analytical_noisy_j1_nm": ("Torque of joint 1 from the analytical inverse dynamics for the same noisy "
                                           "inputs (circles in the frame)", "Nm"),
                "noise_free_j2_nm": ("Noise-free torque of joint 2 (black line)", "Nm"),
                "annet_noisy_j2_nm": ("ANNet torque of joint 2 for the noisy inputs (coloured trace)", "Nm"),
                "analytical_noisy_j2_nm": ("Torque of joint 2 from the analytical inverse dynamics for the same noisy "
                                           "inputs (circles in the frame)", "Nm"),
            },
        ),
        Sheet(
            name="Fig. 7b-c", member="fig7_rmse_values.csv", panel="b, c",
            content="Torque RMSE of every test trajectory under every noise condition. b draws the ANNet rows with "
                    "noise_target none or all (dots and thin lines), their medians (heavy line) and the medians of "
                    "the analytical rows with noise_target all (circles). c draws the ANNet rows with noise_target "
                    "none and the ANNet rows with noise_percent 5. The remaining rows enter the statistics sheets.",
            n="n = 50 test trajectories per model, noise condition, noise level and joint, one noise realization "
              "each",
            shape=(4100, 7),
            columns={
                "model": ("Model, ANNet or the analytical inverse dynamics", NONE),
                "noise_target": ("Inputs that receive the noise, none, all, angle, velocity or acceleration", NONE),
                "noise_percent": NOISE_PERCENT_INPUT,
                "joint": JOINT["joint"],
                **TRAJECTORY,
                "rmse_nm": ("Torque RMSE relative to the noise-free torques (y axis)", "Nm"),
                "torque_sd_nm": ("Standard deviation of the noise-free torque over the trajectory", "Nm"),
            },
        ),
        Sheet(
            name="Fig. 7b statistics", member="fig7_statistics.csv", panel="b",
            content="ANNet against the analytical inverse dynamics on the same noisy inputs (noise on all inputs), "
                    "for each joint and noise level.",
            n="n = 50 trajectories per group, 10 comparisons",
            shape=(10, 30),
            columns={
                **JOINT,
                "noise_percent": NOISE_PERCENT_INPUT,
                "n_trajectories": ("Number of trajectories per group", COUNT),
                "shapiro_p_annet": ("Shapiro-Wilk p value of the ANNet errors", NONE),
                "shapiro_p_analytical": ("Shapiro-Wilk p value of the analytical errors", NONE),
                **primary_test("ANNet", "analytical", "across the 10 comparisons"),
                "mann_whitney_u_annet_vs_analytical": ("U statistic of the two-sided Mann-Whitney U test, ANNet "
                                                       "against analytical", NONE),
                "mann_whitney_p_two_sided": ("p value of the Mann-Whitney U test", NONE),
                **paired_test("ANNet", "analytical"),
                "median_annet_nm": ("Median ANNet error (heavy line)", "Nm"),
                "iqr_annet_nm": ("Interquartile range of the ANNet errors", "Nm"),
                "median_analytical_nm": ("Median analytical error (circle)", "Nm"),
                "iqr_analytical_nm": ("Interquartile range of the analytical errors", "Nm"),
                "median_ratio_annet_over_analytical": ("Median ANNet error divided by the median analytical error",
                                                       NONE),
                "cliffs_delta_annet_vs_analytical": ("Cliff's delta, ANNet against analytical", NONE),
                "n_annet_higher": ("Number of trajectories on which the ANNet error is the higher one", COUNT),
                "median_clean_annet_nm": ("Median ANNet error without noise", "Nm"),
                "median_ratio_noisy_over_clean_annet": ("Median ANNet error with noise divided by that without noise",
                                                        NONE),
                "median_annet_percent_of_torque_sd": ("Median over trajectories of the ANNet error in percent of the "
                                                      "s.d. of the noise-free torque", "%"),
                "median_added_by_annet_nm": ("Median over trajectories of the square root of the squared ANNet error "
                                             "minus the squared analytical error, zero where the analytical error is "
                                             "the larger", "Nm"),
                "median_annet_angle_noise_only_nm": ("Median ANNet error at this level with noise on the angles alone",
                                                     "Nm"),
                "median_annet_velocity_noise_only_nm": ("Median ANNet error at this level with noise on the "
                                                        "velocities alone", "Nm"),
                "median_annet_acceleration_noise_only_nm": ("Median ANNet error at this level with noise on the "
                                                            "accelerations alone", "Nm"),
            },
        ),
        Sheet(
            name="Fig. 7b scaling law", member="fig7_scaling.csv", panel="b",
            content="Growth of the torque error with the noise level for each trajectory, compared with first-order "
                    "propagation of the noise through the inverse dynamics. These values are not drawn. They "
                    "underlie the scaling results reported with Fig. 7b.",
            n="n = 50 test trajectories per joint",
            shape=(100, 10),
            columns={
                "joint": JOINT["joint"],
                **TRAJECTORY,
                "clean_annet_nm": ("ANNet error without noise", "Nm"),
                "gain_first_order_nm_per_percent": ("First-order gain of the torque error from the Jacobian of the "
                                                    "analytical inverse dynamics", "Nm per % of noise"),
                "gain_annet_jacobian_nm_per_percent": ("The same gain from the Jacobian of the ANNet torque",
                                                       "Nm per % of noise"),
                "gain_fitted_analytical_nm_per_percent": ("Least-squares slope through the origin of the analytical "
                                                          "error against the noise level", "Nm per % of noise"),
                "exponent_analytical": ("Slope of the logarithm of the analytical error against the logarithm of the "
                                        "noise level", NONE),
                "exponent_annet": ("Slope of the logarithm of the ANNet error against the logarithm of the noise "
                                   "level", NONE),
                "max_relative_deviation_analytical_from_proportional_law": (
                    "Largest relative deviation of the analytical errors from the product of gain and noise level",
                    NONE),
                "max_relative_deviation_annet_from_quadrature_law": (
                    "Largest relative deviation of the ANNet errors from the square root of the squared noise-free "
                    "error plus the squared product of gain and noise level", NONE),
            },
        ),
        Sheet(
            name="Fig. 7c statistics", member="fig7_input_group_statistics.csv", panel="c",
            content="ANNet errors with noise on one group of inputs or on all inputs against the noise-free ANNet "
                    "errors of the same trajectories. The asterisks of c follow the rows with noise_percent 5.",
            n="n = 50 trajectories per group, 40 comparisons",
            shape=(40, 20),
            columns={
                **JOINT,
                "noise_percent": NOISE_PERCENT_INPUT,
                "noise_target": ("Inputs that receive the noise, angle, velocity, acceleration or all", NONE),
                "n_trajectories": ("Number of trajectories per group", COUNT),
                "shapiro_p_noisy": ("Shapiro-Wilk p value of the errors with noise", NONE),
                "shapiro_p_noise_free": ("Shapiro-Wilk p value of the errors without noise", NONE),
                **primary_test("with noise", "without noise", "across the 40 comparisons"),
                **paired_test("with noise", "without noise"),
                "median_noise_free_annet_nm": ("Median ANNet error without noise", "Nm"),
                "median_annet_nm": ("Median ANNet error with noise", "Nm"),
                "median_ratio_noisy_over_noise_free": ("Median error with noise divided by that without noise", NONE),
                "cliffs_delta_noisy_vs_noise_free": ("Cliff's delta, with noise against without noise", NONE),
                "n_higher_than_noise_free": ("Number of trajectories whose error is higher with noise", COUNT),
            },
        ),
        Sheet(
            name="Fig. 7d", member="fig7_filter_rmse.csv", panel="d",
            content="Torque RMSE of every test trajectory after zero-phase low-pass filtering of the noisy inputs "
                    "(noise on all inputs). d draws the rows with noise_percent 5, the ANNet rows as dots and thin "
                    "lines with their medians as the heavy line and the medians of the analytical rows as circles. "
                    "The dashed and dotted lines are the ANNet medians at 5% noise and without noise on the sheet "
                    "Fig. 7b-c.",
            n="n = 50 test trajectories per model, noise level, cutoff and joint",
            shape=(12000, 6),
            columns={
                "model": ("Model, ANNet or the analytical inverse dynamics", NONE),
                "noise_percent": ("Noise level as on the sheet Fig. 7b-c, 0 for noise-free inputs", "%"),
                "cutoff_hz": ("Cutoff frequency of the low-pass filter (x axis)", "Hz"),
                "joint": JOINT["joint"],
                **TRAJECTORY,
                "rmse_nm": ("Torque RMSE relative to the noise-free torques (y axis)", "Nm"),
            },
        ),
        Sheet(
            name="Fig. 7d statistics", member="fig7_filter_statistics.csv", panel="d",
            content="Effect of the filter for each joint, noise level and cutoff, with the test of filtered against "
                    "unfiltered ANNet errors.",
            n="n = 50 trajectories per group, 120 combinations of joint, noise level and cutoff",
            shape=(120, 16),
            columns={
                **JOINT,
                "noise_percent": ("Noise level as on the sheet Fig. 7b-c, 0 for noise-free inputs", "%"),
                "cutoff_hz": ("Cutoff frequency of the low-pass filter", "Hz"),
                "n_trajectories": ("Number of trajectories", COUNT),
                "white_noise_sd_fraction": ("Fraction of the s.d. of white noise that the filter transmits", NONE),
                "median_unfiltered_annet_nm": ("Median ANNet error without filtering", "Nm"),
                "median_filtered_annet_nm": ("Median ANNet error with filtering", "Nm"),
                "median_filtered_analytical_nm": ("Median analytical error with filtering", "Nm"),
                "median_ratio_annet_over_analytical": ("median_filtered_annet_nm divided by "
                                                       "median_filtered_analytical_nm", NONE),
                "median_filter_only_analytical_nm": ("Median analytical error that the filter alone causes on "
                                                     "noise-free inputs", "Nm"),
                "reduction_factor_of_median_annet": ("median_unfiltered_annet_nm divided by median_filtered_annet_nm",
                                                     NONE),
                "n_filtered_lower_than_unfiltered": ("Number of trajectories whose ANNet error is lower with "
                                                     "filtering", COUNT),
                "wilcoxon_filtered_vs_unfiltered_p_two_sided": ("p value of the two-sided Wilcoxon signed-rank test "
                                                                "of filtered against unfiltered ANNet errors, paired "
                                                                "by trajectory", NONE),
                "median_analytical_over_first_order_prediction": ("Median ratio of the analytical error to its "
                                                                  "first-order prediction", NONE),
                "lowest_median_annet_of_level": ("True for the cutoff with the lowest median ANNet error at that "
                                                 "noise level", NONE),
            },
        ),
    ],
)

# ----------------------------------------------------------------------------------
# Fig. 8
# ----------------------------------------------------------------------------------
NOISE_PERCENT_TORQUE = ("Noise level, the s.d. of the noise in percent of the s.d. of each gravitational torque over "
                        "the noise-free trajectory", "%")
SOLVER_QUANTITIES = {"g": ("Gravitational torque", "Nm"), "q": ("Angle", "rad"), "dq": ("Velocity", "rad/s"),
                     "ddq": ("Acceleration", "rad/s^2")}
FIG8_LEVELS = ("0", "1", "2.5", "5", "10")
FIG8_SOURCES = {"reference": "noise-free fourth-order Runge-Kutta reference"}
for _solver, _text in (("annet", "solver with the learned energy (ANNet)"),
                       ("analytical", "solver with the analytical energy")):
    for _level in FIG8_LEVELS:
        FIG8_SOURCES[f"{_solver}_{_level}pct"] = f"{_text} at {_level}% torque noise"
FIG8_DRAWN = {(source, quantity) for source in ("reference", "annet_10pct") for quantity in SOLVER_QUANTITIES}


def level_statistics(group_a, group_b):
    return {
        "shapiro_p_a": (f"Shapiro-Wilk p value of the errors {group_a}", NONE),
        "shapiro_p_b": (f"Shapiro-Wilk p value of the errors {group_b}", NONE),
    }


FIG8 = Figure(
    label="Fig. 8", archive="FIG8.zip",
    title="The learned energy preserves the tolerance of the forward solver to noise",
    sheets=[
        Sheet(
            name="Fig. 8a", member="fig8_example_traces.csv", panel="a",
            content="Example trajectory at every solver step, for the reference and for both solvers at every noise "
                    "level. a draws the reference columns and the columns of the solver with the learned energy at "
                    "10% noise. For a solver, g is the gravitational torque that entered it and q, dq and ddq are "
                    "the angle, velocity and acceleration that it returned.",
            n="n = 1 trajectory (index 44 of the 50 trajectories), 2,500 time points",
            shape=(2500, 89),
            columns=trace_columns(FIG8_SOURCES, SOLVER_QUANTITIES, FIG8_DRAWN),
        ),
        Sheet(
            name="Fig. 8b", member="fig8_rmse_values.csv", panel="b",
            content="RMSE of acceleration, velocity and angle of every trajectory at every noise level. b draws the "
                    "ANNet rows (dots and thin lines), their medians (heavy lines) and the medians of the analytical "
                    "rows (circles).",
            n="n = 50 trajectories per solver, noise level, quantity and joint, one noise realization each",
            shape=(3000, 8),
            columns={
                "model": ("Energy in the forward solver, ANNet (learned) or analytical", NONE),
                "noise_percent": NOISE_PERCENT_TORQUE,
                **QUANTITY,
                "joint": JOINT["joint"],
                **TRAJECTORY,
                "rmse": ("RMSE relative to the noise-free reference (y axis)", "see column unit"),
                "torque_sd_nm": ("Standard deviation of the gravitational torque of the joint over the noise-free "
                                 "reference trajectory", "Nm"),
            },
        ),
        Sheet(
            name="Fig. 8b level statistics", member="fig8_level_statistics.csv", panel="b",
            content="Errors at each noise level against the noise-free errors of the same solver. The asterisks and "
                    "daggers of b follow the rows with model ANNet.",
            n="n = 50 trajectories per group, 24 comparisons per solver",
            shape=(48, 27),
            columns={
                "model": ("Energy in the forward solver, ANNet (learned) or analytical", NONE),
                **QUANTITY,
                **JOINT,
                "noise_percent": NOISE_PERCENT_TORQUE,
                "n_trajectories": ("Number of trajectories per group", COUNT),
                **level_statistics("with noise", "without noise"),
                **primary_test("with noise", "without noise", "across the 24 comparisons of the solver"),
                **paired_test("with noise", "without noise"),
                "median_noise_free": ("Median error without noise", "see column unit"),
                "median_noisy": ("Median error with noise", "see column unit"),
                "median_ratio_noisy_over_noise_free": ("Median error with noise divided by that without noise", NONE),
                "median_paired_difference_noisy_minus_noise_free": ("Median over trajectories of the error with noise "
                                                                    "minus the error without noise",
                                                                    "see column unit"),
                "direction": ("higher or lower, the median error with noise relative to that without noise", NONE),
                "cliffs_delta_noisy_vs_noise_free": ("Cliff's delta, with noise against without noise", NONE),
                "n_higher_than_noise_free": ("Number of trajectories whose error is higher with noise", COUNT),
                "median_added_by_noise": ("Median over trajectories of the square root of the squared error with "
                                          "noise minus the squared error without noise, zero where the error with "
                                          "noise is the smaller", "see column unit"),
                "median_added_by_noise_per_percent": ("median_added_by_noise divided by the noise level",
                                                      "column unit per %"),
                "wilcoxon_paired_fdr_q_value": ("Benjamini-Hochberg adjusted p value of the paired Wilcoxon test, "
                                                "across the 24 comparisons of the solver", NONE),
            },
        ),
        Sheet(
            name="Fig. 8b solver statistics", member="fig8_solver_statistics.csv", panel="b",
            content="Errors of the solver with the learned energy (ANNet) against those of the solver with the "
                    "analytical energy under the same noise.",
            n="n = 50 trajectories per group, 30 comparisons",
            shape=(30, 23),
            columns={
                **QUANTITY,
                **JOINT,
                "noise_percent": NOISE_PERCENT_TORQUE,
                "n_trajectories": ("Number of trajectories per group", COUNT),
                **level_statistics("of ANNet", "of the analytical energy"),
                **primary_test("ANNet", "analytical", "across the 30 comparisons"),
                **paired_test("ANNet", "analytical"),
                "median_annet": ("Median error of ANNet (heavy line)", "see column unit"),
                "median_analytical": ("Median error of the analytical energy (circle)", "see column unit"),
                "median_ratio_annet_over_analytical": ("Median error of ANNet divided by that of the analytical "
                                                       "energy", NONE),
                "median_paired_difference_annet_minus_analytical": ("Median over trajectories of the ANNet error "
                                                                    "minus the analytical error", "see column unit"),
                "cliffs_delta_annet_vs_analytical": ("Cliff's delta, ANNet against analytical", NONE),
                "n_annet_higher": ("Number of trajectories on which the ANNet error is the higher one", COUNT),
                "wilcoxon_paired_fdr_q_value": ("Benjamini-Hochberg adjusted p value of the paired Wilcoxon test, "
                                                "across the 30 comparisons", NONE),
            },
        ),
    ],
)

# ----------------------------------------------------------------------------------
# Fig. 9
# ----------------------------------------------------------------------------------
FIG9_SOURCES = {
    "reference": "fourth-order Runge-Kutta reference of the perturbed pendulum",
    "before": "forward solver with the network before adaptation",
    "after": "forward solver with the network after adaptation",
    "analytical": "forward solver with the analytical acceleration energy",
}
FIG9_QUANTITIES = {key: SOLVER_QUANTITIES[key] for key in ("q", "dq", "ddq")}
FIG9_DRAWN = {(source, quantity) for source in ("reference", "before", "after") for quantity in ("q", "dq")}
PLANT = ("Simulated pendulum, named by its changed parameters (unchanged, mass, length or mass_length)", NONE)

FIG9 = Figure(
    label="Fig. 9", archive="FIG9.zip",
    title="Learning of the inverse dynamics transfers to the forward dynamics without forward-specific training",
    sheets=[
        Sheet(
            name="Fig. 9a", member="fig9_example_traces.csv", panel="a",
            content="Example trajectory of the pendulum with changed masses and link lengths at every solver step. "
                    "a draws the angle (q) and the velocity (dq) of the reference and of the solver before and "
                    "after adaptation.",
            n="n = 1 trajectory (index 10 of the 50 trajectories), 2,500 time points",
            shape=(2500, 25),
            columns=trace_columns(FIG9_SOURCES, FIG9_QUANTITIES, FIG9_DRAWN),
        ),
        Sheet(
            name="Fig. 9b-c", member="fig9_rmse_values.csv", panel="b, c",
            content="RMSE of every trajectory in every condition. b draws the rows of plant mass_length with energy "
                    "network and quantity torque or angle at epochs 0 to 300 (dots and thin lines), their medians "
                    "(heavy lines) and the medians of the same rows of plant unchanged (dashed lines). c draws the "
                    "rows with task forward, quantity angle and stage before, after or analytical for the four "
                    "pendulums. The remaining rows enter the statistics sheet.",
            n="n = 50 trajectories per condition and joint",
            shape=(13300, 10),
            columns={
                "task": ("forward for the closed-loop forward simulation, inverse for the torque prediction on the "
                         "reference trajectories", NONE),
                "plant": PLANT,
                "energy": ("Acceleration energy in use, network or analytical", NONE),
                "epoch": ("Adaptation epoch of the network (x axis of b), empty for the analytical energy and for "
                          "the network that is not updated", NONE),
                "stage": ("before (epoch 0), intermediate, after (epoch 300), analytical, or not_updated (pretrained "
                          "network with the gravitational torque of the original pendulum)", NONE),
                "quantity": ("Quantity that the error refers to, angle, velocity, acceleration or torque", NONE),
                "unit": ("Unit of the error of that quantity", NONE),
                "joint": JOINT["joint"],
                **TRAJECTORY,
                "rmse": ("RMSE relative to the reference of the simulated pendulum (y axis)", "see column unit"),
            },
        ),
        Sheet(
            name="Fig. 9b-c statistics", member="fig9_statistics.csv", panel="b, c",
            content="Comparisons of group a with group b in 8 families, with the p values adjusted within each "
                    "family. The asterisks of b follow the families torque_epoch_vs_before (36 comparisons) and "
                    "epoch_vs_before (108) for plant mass_length. The brackets of c follow, for quantity angle, the "
                    "families epoch_vs_before and control_epoch_vs_before (36) at epoch 300 and after_vs_analytical "
                    "(24). The other families are control_torque_epoch_vs_before (12), after_vs_control (18), "
                    "after_vs_original (18) and after_vs_not_updated (18).",
            n="n = 50 trajectories per group, 270 comparisons",
            shape=(270, 29),
            columns={
                "family": ("Family of comparisons within which the p values are adjusted", NONE),
                "plant": PLANT,
                "quantity": ("Quantity that the error refers to, angle, velocity, acceleration or torque", NONE),
                "unit": ("Unit of the error of that quantity", NONE),
                **JOINT,
                "group_a": ("First group of the comparison", NONE),
                "group_b": ("Second group of the comparison", NONE),
                "epoch_a": ("Adaptation epoch of the network of group a", NONE),
                "n_trajectories": ("Number of trajectories per group", COUNT),
                **level_statistics("of group a", "of group b"),
                **primary_test("group a", "group b", "within the family"),
                **paired_test("group a", "group b"),
                "median_a": ("Median error of group a", "see column unit"),
                "median_b": ("Median error of group b", "see column unit"),
                "median_ratio_a_over_b": ("Median error of group a divided by that of group b", NONE),
                "median_paired_difference_a_minus_b": ("Median over trajectories of the error of group a minus that "
                                                       "of group b", "see column unit"),
                "direction": ("higher or lower, the median of group a relative to that of group b", NONE),
                "cliffs_delta_a_vs_b": ("Cliff's delta, group a against group b", NONE),
                "n_a_lower": ("Number of trajectories on which group a has the lower error", COUNT),
                "wilcoxon_paired_fdr_q_value": ("Benjamini-Hochberg adjusted p value of the paired Wilcoxon test, "
                                                "within the family", NONE),
                "wilcoxon_paired_significant_fdr_0_05": ("True when the adjusted p value of the paired Wilcoxon test "
                                                         "is below 0.05", NONE),
            },
        ),
    ],
)

# ----------------------------------------------------------------------------------
# Supplementary Fig. 1
# ----------------------------------------------------------------------------------
FIGS1 = Figure(
    label="Supplementary Fig. 1", archive="FIGS1.zip",
    title="Parallel-coordinate representation of hyperparameter optimization",
    sheets=[
        Sheet(
            name="Supplementary Fig. 1", member="figs1_parallel_coordinates_trials.csv", panel="single panel",
            content="One row per trial of the hyperparameter search, drawn as one polyline. The axes are, from left "
                    "to right, params_lr (logarithmic), activation_code, params_n_layers, params_n_units_l0 to "
                    "params_n_units_l9 and RMSE (logarithmic).",
            n="n = 500 trials",
            shape=(500, 21),
            columns={
                "number": ("Index of the trial in the search", NONE),
                "value": ("Objective value of the trial, the validation mean squared torque error", "Nm^2"),
                "RMSE": ("Validation torque RMSE, the square root of value (last axis)", "Nm"),
                "state": ("Final state of the trial, COMPLETE or PRUNED", NONE),
                "params_lr": ("Learning rate (first axis)", NONE),
                "params_activation": ("Activation function", NONE),
                "activation_code": ("Position of the activation function on its axis, 0 Tanh, 1 Softplus, 2 ReLU, "
                                    "3 ELU, 4 GELU", NONE),
                "params_n_layers": ("Number of hidden layers (depth axis)", COUNT),
                **{f"params_n_units_l{k}": (f"Nodes in hidden layer {k + 1}, 0 when the layer is absent", COUNT)
                   for k in range(10)},
                "top_decile_rmse": ("True for the 50 trials with the lowest RMSE (teal lines)", NONE),
                "best_trial": ("True for the best trial (red line)", NONE),
                "unused_deeper_layers_encoded_as_zero": ("True when the network has fewer than ten hidden layers, so "
                                                         "that the widths of the absent layers are written as 0",
                                                         NONE),
            },
        ),
    ],
)

# ----------------------------------------------------------------------------------
# Supplementary Fig. 2
# ----------------------------------------------------------------------------------
SIZES = (1000, 2000, 4000, 8000, 16000, 32000, 64000, 128000, 256000)


def q_matrix(letter, joint):
    return Sheet(
        name=f"Supplementary Fig. 2{letter}", member=f"figs2_fdr_q_matrix_J{joint}.csv", panel=letter,
        content=f"Matrix of adjusted p values of the pairwise comparisons for joint {joint} (cell colour). Rows and "
                "columns are training-set sizes, and the diagonal is empty.",
        n="n = 50 test trajectories per training-set size, 36 comparisons",
        shape=(9, 10),
        columns={
            "dataset_size": ("Training-set size of the row", COUNT),
            **{str(size): (f"Benjamini-Hochberg adjusted p value of the comparison with the network trained on "
                           f"{size:,} samples", NONE) for size in SIZES},
        },
    )


FIGS2 = Figure(
    label="Supplementary Fig. 2", archive="FIGS2.zip",
    title="Pairwise statistical comparison of inverse-dynamics accuracy across training-set sizes",
    note="The torque RMSE values compared by these tests are in Supplementary Data 2 (source data for Fig. 3), sheet Fig. 3d upper.",
    sheets=[
        q_matrix("a", 1),
        q_matrix("b", 2),
        Sheet(
            name="Supplementary Fig. 2 statistics", member="figs2_pairwise_statistics.csv", panel="a, b",
            content="The pairwise comparisons behind both matrices, one row per pair of training-set sizes and "
                    "joint.",
            n="n = 50 test trajectories per training-set size, 36 comparisons per joint",
            shape=(72, 13),
            columns={
                "joint": JOINT["joint"],
                "dataset_size_a": ("Training-set size of the first network of the pair", COUNT),
                "dataset_label_a": ("Label of that size in the figure", NONE),
                "dataset_size_b": ("Training-set size of the second network of the pair", COUNT),
                "dataset_label_b": ("Label of that size in the figure", NONE),
                "n_a": ("Number of test trajectories of the first network", COUNT),
                "n_b": ("Number of test trajectories of the second network", COUNT),
                "median_rmse_a_nm": ("Median torque RMSE of the first network", "Nm"),
                "median_rmse_b_nm": ("Median torque RMSE of the second network", "Nm"),
                "mann_whitney_u": ("U statistic of the two-sided Mann-Whitney U test", NONE),
                "raw_p_value": ("p value of that test", NONE),
                "fdr_q_value": ("Benjamini-Hochberg adjusted p value across the 36 comparisons of the joint", NONE),
                "significant_fdr_0_05": ("True when the adjusted p value is below 0.05", NONE),
            },
        ),
    ],
)

# ----------------------------------------------------------------------------------
# Supplementary Fig. 3
# ----------------------------------------------------------------------------------
FIGS3 = Figure(
    label="Supplementary Fig. 3", archive="FIGS3.zip",
    title="Learned and analytical acceleration energies give median closed-loop simulation errors within 8% of each "
          "other",
    sheets=[
        Sheet(
            name="Supplementary Fig. 3a-f", member="figs3_rmse_values.csv", panel="a-f",
            content="RMSE of every trajectory of the closed-loop forward simulation with the learned and with the "
                    "analytical acceleration energy. Joint 1 is in a (angle, quantity position), b (velocity) and c "
                    "(acceleration), joint 2 in d, e and f. A line joins the two rows that share quantity, joint "
                    "and trajectory.",
            n="n = 50 trajectories per energy, quantity and joint",
            shape=(600, 6),
            columns={
                "model": ("Energy in the forward solver, learned_S (learned, filled circles) or analytical_S "
                          "(analytical Gibbs-Appell function, open circles)", NONE),
                "quantity": ("Quantity that the error refers to, position (angle), velocity or acceleration", NONE),
                "unit": ("Unit of the error of that quantity", NONE),
                "joint": JOINT["joint"],
                **TRAJECTORY,
                "rmse": ("RMSE relative to the fourth-order Runge-Kutta reference (y axis)", "see column unit"),
            },
        ),
        Sheet(
            name="Supplementary Fig. 3 statistics", member="figs3_statistics.csv", panel="a-f",
            content="Comparison of the two energies for each quantity and joint.",
            n="n = 50 trajectories per group, 6 comparisons",
            shape=(6, 28),
            columns={
                "quantity": ("Quantity that the error refers to, position (angle), velocity or acceleration", NONE),
                "unit": ("Unit of the error of that quantity", NONE),
                **JOINT,
                "n_learned": ("Number of trajectories with the learned energy", COUNT),
                "n_analytical": ("Number of trajectories with the analytical energy", COUNT),
                "shapiro_p_learned": ("Shapiro-Wilk p value of the errors with the learned energy", NONE),
                "shapiro_p_analytical": ("Shapiro-Wilk p value of the errors with the analytical energy", NONE),
                **primary_test("analytical", "learned", "across the 6 comparisons"),
                "mann_whitney_u_analytical_vs_learned": ("U statistic of the two-sided Mann-Whitney U test, "
                                                         "analytical against learned", NONE),
                "mann_whitney_p_two_sided": ("p value of the Mann-Whitney U test", NONE),
                **paired_test("analytical", "learned"),
                "mean_learned": ("Mean error with the learned energy", "see column unit"),
                "sd_learned": ("Standard deviation of the errors with the learned energy", "see column unit"),
                "median_learned": ("Median error with the learned energy (heavy bar)", "see column unit"),
                "iqr_learned": ("Interquartile range of the errors with the learned energy", "see column unit"),
                "mean_analytical": ("Mean error with the analytical energy", "see column unit"),
                "sd_analytical": ("Standard deviation of the errors with the analytical energy", "see column unit"),
                "median_analytical": ("Median error with the analytical energy (heavy bar)", "see column unit"),
                "iqr_analytical": ("Interquartile range of the errors with the analytical energy", "see column unit"),
                "median_ratio_analytical_over_learned": ("Median error with the analytical energy divided by that "
                                                         "with the learned energy", NONE),
                "cliffs_delta_analytical_vs_learned": ("Cliff's delta, analytical against learned", NONE),
            },
        ),
    ],
)

FIGURES = [FIG2, FIG3, FIG4, FIG5, FIG6, FIG7, FIG8, FIG9, FIGS1, FIGS2, FIGS3]


# ----------------------------------------------------------------------------------
# Reading, writing and checking
# ----------------------------------------------------------------------------------
def read_member(archive, member):
    """Table of an archive as a data frame, with every number parsed to the nearest double."""
    data = archive.read(member)
    if member.endswith(".json"):
        return pd.DataFrame([json.loads(data)])
    return pd.read_csv(io.BytesIO(data), float_precision="round_trip", keep_default_na=False, na_values=[""])


def digest(table):
    """SHA-256 of the numeric values of a table (as doubles, column by column) and of its text values."""
    numbers, text = hashlib.sha256(), hashlib.sha256()
    for name in table.columns:
        column = table[name]
        if pd.api.types.is_bool_dtype(column) or pd.api.types.is_numeric_dtype(column):
            values = column.to_numpy(dtype=np.float64)
            values = np.where(np.isnan(values), np.nan, values)   # one bit pattern for every empty cell
            numbers.update(values.astype("<f8").tobytes())
        else:
            text.update("\x1f".join("" if pd.isna(value) else str(value) for value in column).encode("utf-8"))
        text.update(str(name).encode("utf-8"))
    return numbers.hexdigest(), text.hexdigest()


def write_readme(book, figure, tables):
    readme = book.create_sheet("README", 0)
    bold = Font(bold=True)
    readme.append([f"{figure.name}. Source data for {figure.label}"])
    readme.append([figure.title])
    readme.append([f"Each sheet holds one table of the figure data archive {figure.archive}, with unrounded values "
                   "in the order of the archive. Empty cells mark entries that do not apply."])
    if figure.note:
        readme.append([figure.note])
    readme.append([])
    readme.append(["Sheet", "Panel", "Content", "n", "Rows", "Archive table"])
    for cell in readme[readme.max_row]:
        cell.font = bold
    for sheet in figure.sheets:
        readme.append([sheet.name, sheet.panel, sheet.content, sheet.n, len(tables[sheet.name]), sheet.member])
    readme.append([])
    readme.append(["Sheet", "Column", "Description", "Unit"])
    for cell in readme[readme.max_row]:
        cell.font = bold
    for sheet in figure.sheets:
        for column in tables[sheet.name].columns:
            description, unit = sheet.columns[str(column)]
            readme.append([sheet.name, str(column), description, unit])
    book.active = 0


CORE_PROPERTIES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>{title}</dc:title></cp:coreProperties>'
)


def strip_package_metadata(path, title):
    """Rewrite the workbook package without author or dates, in its properties and on its entries."""
    with zipfile.ZipFile(path) as package:
        entries = [(info.filename, package.read(info.filename)) for info in package.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in entries:
            if name == "docProps/core.xml":
                data = CORE_PROPERTIES.format(title=title).encode("utf-8")
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            package.writestr(info, data)


@contextmanager
def exact_numbers():
    """Store every number with the digits that restore its double exactly.

    openpyxl writes numbers with 16 significant digits, one fewer than some doubles need.
    """
    default = cell_writer.safe_string

    def safe_string(value):
        if isinstance(value, float) and math.isfinite(value):
            return repr(float(value))
        return default(value)

    cell_writer.safe_string = safe_string
    try:
        yield
    finally:
        cell_writer.safe_string = default


def build_workbook(figure, archive_dir, output_dir):
    with zipfile.ZipFile(archive_dir / figure.archive) as archive:
        tables = {sheet.name: read_member(archive, sheet.member) for sheet in figure.sheets}

    for sheet in figure.sheets:
        table = tables[sheet.name]
        assert len(sheet.name) <= 31, sheet.name
        assert table.shape == sheet.shape, (sheet.name, table.shape, sheet.shape)
        assert len(table) < EXCEL_MAX_ROWS, sheet.name
        missing = [str(column) for column in table.columns if str(column) not in sheet.columns]
        unused = [column for column in sheet.columns if column not in map(str, table.columns)]
        assert not missing and not unused, (sheet.name, missing, unused)
        numeric = table.select_dtypes(include="number").to_numpy(dtype=float)
        assert not np.isinf(numeric).any(), sheet.name

    path = output_dir / figure.workbook
    with exact_numbers(), pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet in figure.sheets:
            tables[sheet.name].to_excel(writer, sheet_name=sheet.name, index=False, freeze_panes=(1, 0))
            for cell in writer.sheets[sheet.name][1]:
                cell.font, cell.border, cell.alignment = Font(bold=True), Border(), Alignment()
        write_readme(writer.book, figure, tables)
    strip_package_metadata(path, f"{figure.name}. Source data for {figure.label}")
    return path, tables


def check_workbook(figure, path, tables):
    """Read the workbook back and compare every sheet with the table of the archive."""
    report = []
    for sheet in figure.sheets:
        source = tables[sheet.name]
        stored = pd.read_excel(path, sheet_name=sheet.name, engine="openpyxl", keep_default_na=False, na_values=[""])
        same_shape = stored.shape == source.shape
        same_header = [str(column) for column in stored.columns] == [str(column) for column in source.columns]
        source_digest, stored_digest = digest(source), digest(stored)
        numeric_cells = int(sum(source[name].notna().sum() for name in source.columns
                                if pd.api.types.is_numeric_dtype(source[name])))
        report.append({
            "workbook": path.name, "sheet": sheet.name, "member": sheet.member,
            "rows": len(source), "columns": source.shape[1], "numeric_cells": numeric_cells,
            "numeric_sha256": source_digest[0][:16],
            "match": bool(same_shape and same_header and source_digest == stored_digest),
        })
    return report


def main():
    here = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="Build the source-data workbooks from the figure archives.")
    parser.add_argument("output_dir", nargs="?", type=Path, default=here,
                        help="folder for the workbooks (default: the folder of this script)")
    parser.add_argument("--archives", type=Path, default=here.parent / "FIGURES" / "DATA AND SCRIPTS",
                        help="folder that holds FIG2.zip ... FIGS3.zip (default: ../FIGURES/DATA AND SCRIPTS)")
    arguments = parser.parse_args()
    arguments.output_dir.mkdir(parents=True, exist_ok=True)

    report = []
    for figure in FIGURES:
        path, tables = build_workbook(figure, arguments.archives, arguments.output_dir)
        report.extend(check_workbook(figure, path, tables))
    report = pd.DataFrame(report)
    print(report.to_string(index=False))
    print(f"\n{report['workbook'].nunique()} workbooks, {len(report)} data sheets, "
          f"{int(report['rows'].sum()):,} rows, {int(report['match'].sum())} of {len(report)} sheets identical "
          "to the archive.")
    return 0 if report["match"].all() else 1


if __name__ == "__main__":
    sys.exit(main())
